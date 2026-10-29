"""Immutable exam snapshots, escaped HTML/structured SVG, editable Word equations."""
import html
import json
from pathlib import Path
import re
from .validation import printable, validate_content, validate_diagram

ASSETS = Path(__file__).parent / "assets"


def svg(objects):
    validate_diagram(objects)
    if not objects:
        return ""
    result = ['<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1000 1000" width="280" height="280" role="img">']
    for item in objects:
        kind = item["type"]
        if kind == "text":
            result.append(f'<text x="{item["x"]}" y="{item["y"]}" font-size="38">{html.escape(item["label"])}</text>')
        elif kind == "polyline":
            points = " ".join(f"{x},{y}" for x, y in item["points"])
            result.append(f'<polyline points="{points}" fill="none" stroke="#142c40" stroke-width="3"/>')
        else:
            keys = ("x1", "y1", "x2", "y2") if kind == "line" else ("cx", "cy", "r")
            attrs = " ".join(f'{key}="{item[key]}"' for key in keys)
            result.append(f'<{kind} {attrs} fill="none" stroke="#142c40" stroke-width="3"/>')
    return "".join(result) + "</svg>"


def text_html(text):
    return html.escape(str(text)).replace("\n", "<br>")


def page_html(body, title="MathBank Desktop"):
    return '''<!doctype html><html lang="ko"><head><meta charset="utf-8">
<meta http-equiv="Content-Security-Policy" content="default-src 'none'; script-src 'self' file: 'unsafe-inline'; style-src 'self' file: 'unsafe-inline'; font-src 'self' file:; img-src data: file:; connect-src 'none'">
<link rel="stylesheet" href="katex/katex.min.css"><script src="katex/katex.min.js"></script>
<script src="katex/contrib/auto-render.min.js"></script>
<style>
@page {size:A4; margin:17mm 16mm;} *{box-sizing:border-box} body{font-family:'Noto Sans CJK KR','Malgun Gothic',sans-serif;color:#142c40;font-size:11pt;line-height:1.65;margin:0;background:white}
h1{font-size:24pt;letter-spacing:-1px;margin:0 0 8mm} h2{font-size:15pt;color:#087e82}
h3{font-size:12pt} .muted{color:#617482;font-size:9pt}.tag{color:#087e82;font-size:9pt}
.question{padding:5mm 0;border-bottom:1px solid #dce4ea;break-inside:avoid}
.question.long{break-inside:auto} .question p{orphans:3;widows:3} .choices{break-inside:avoid}
.columns{column-count:2;column-gap:10mm}.grid{display:grid;grid-template-columns:1fr 1fr;gap:7mm}
.blank{height:18mm}.warning{background:#fff1d9;padding:3mm;color:#8b4d00}
table{width:100%;border-collapse:collapse;font-size:10pt}td,th{padding:2mm;border-bottom:1px solid #dce4ea;text-align:left}tr{break-inside:avoid}
.report-page{break-before:page}.score{font-size:28pt;font-weight:bold;color:#087e82}
svg{max-width:100%;display:block;margin:auto}.katex-display{overflow:visible}.katex{font-size:1.05em}
</style><title>''' + html.escape(title) + '</title></head><body>' + body + '''
<script>window.renderReady=false;window.renderError='';
window.addEventListener('load',async()=>{try{
renderMathInElement(document.body,{delimiters:[{left:'$$',right:'$$',display:true},{left:'$',right:'$',display:false}],throwOnError:true,trust:false,strict:'error'});
await document.fonts.ready; await Promise.all(Array.from(document.images).map(i=>i.complete?Promise.resolve():new Promise((r,j)=>{i.onload=r;i.onerror=j;})));
window.renderReady=true;}catch(e){window.renderError=String(e);}});</script></body></html>'''


def make_exam(store, ids, title, mode="complete", points=5, layout="one"):
    if not ids or len(ids) != len(set(ids)):
        raise ValueError("중복 없이 문항을 선택하세요.")
    if points <= 0 or layout not in ("one", "two", "grid"):
        raise ValueError("배점 또는 조판 설정 오류")
    questions = [store.question(identifier) for identifier in ids]
    blocked = [q for q in questions if not printable(q)]
    if blocked and mode != "draft":
        raise ValueError("출력 예외: " + ", ".join(f'{q["number"] or q["id"][:8]} ({q["error"] or "자동 검사 미완료"})' for q in blocked))
    return store.put_record("exam", {"title": title, "questions": questions, "points": [points] * len(ids),
                                     "layout": layout, "draft": mode == "draft"})


def exam_html(exam, kind="questions"):
    if kind not in ("questions", "answers", "solutions"):
        raise ValueError("출력 종류 오류")
    suffix = {"questions": "문제지", "answers": "빠른답", "solutions": "상세 해설"}[kind]
    body = f'<div class="tag">MATHBANK / {suffix}</div><h1>{text_html(exam["title"])}</h1>'
    if exam.get("draft"):
        body += '<p class="warning">초안 · 자동 검사 미완료 문항은 아래에 개별 표시됩니다.</p>'
    body += '<p class="muted">이름 ____________________　 날짜 ____________________</p>'
    layout = {"one": "", "two": "columns", "grid": "grid"}[exam.get("layout", "one")]
    body += f'<main class="{layout}">'
    for i, q in enumerate(exam["questions"], 1):
        valid = printable(q, kind)
        if not valid and not exam.get("draft"):
            raise ValueError(f"{i}번: {suffix} 자동 검사 미완료입니다.")
        c = q["content"]
        body += f'<section class="question {"long" if len(c.get("text", "")) > 900 else ""}"><b>{i:02d}</b> <span class="muted">[{exam["points"][i-1]}점]</span>'
        if not valid:
            body += f'<p class="warning">처리 예외: {text_html(q["error"] or "자동 검사 미완료")}</p>'
        if kind == "questions":
            body += '<p>' + text_html(c.get("text", "본문 미처리")) + '</p>'
            body += svg(c.get("diagram", []))
            body += '<div class="choices">' + '<br>'.join(f'{n}. {text_html(t)}' for n, t in enumerate(c.get("choices", []), 1)) + '</div><div class="blank"></div>'
        else:
            body += '<p>' + text_html(c.get("answer", "정답 미처리")) + '</p>'
            if kind == "solutions":
                body += ''.join(f'<p>{n}. {text_html(step)}</p>' for n, step in enumerate(c.get("solution", []), 1))
        body += '</section>'
    return page_html(body + '</main>', exam["title"])


def export_docx(exam, destination):
    from docx import Document
    from docx.shared import Mm
    from docx.oxml import parse_xml
    from latex2mathml.converter import convert
    from mathml2omml import convert as omml
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(17)
    document.add_heading(exam["title"], 0)
    if exam.get("draft"):
        document.add_paragraph("초안 · 자동 검사 미완료 문항 포함")
    for i, q in enumerate(exam["questions"], 1):
        if not printable(q) and not exam.get("draft"):
            raise ValueError(f"{i}번: 원문 자동 검사 미완료")
        c = q["content"]
        if c.get("diagram"):
            raise ValueError(f"{i}번: DOCX 벡터 도형 내보내기 미지원. PDF로 출력하세요. 도형을 누락한 파일은 저장하지 않습니다.")
        document.add_heading(f'{i:02d}  [{exam["points"][i-1]}점]', 2)
        if not printable(q):
            document.add_paragraph("처리 예외: " + (q["error"] or "자동 검사 미완료"))
        for text in [c.get("text", "본문 미처리"), *[f"{n}. {t}" for n, t in enumerate(c.get("choices", []), 1)]]:
            paragraph = document.add_paragraph()
            for part in re.split(r'(\$\$.*?\$\$|\$[^$]+\$)', text, flags=re.S):
                if part.startswith('$'):
                    try:
                        equation = omml(convert(part.strip('$')))
                        paragraph._p.append(parse_xml('<root xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">' + equation + '</root>')[0])
                    except Exception as exc:
                        raise ValueError(f"{i}번 수식 변환 실패: {part}") from exc
                else:
                    paragraph.add_run(part)
        document.add_paragraph("\n")
    document.save(destination)


def grade_report(exam, student, grades, notes=None, expected_grade=""):
    if len(grades) != len(exam["questions"]):
        raise ValueError("문항 수와 채점 수가 다릅니다.")
    earned, possible, graded = 0, sum(exam["points"]), 0
    matrix = {}
    for i, (q, grade) in enumerate(zip(exam["questions"], grades)):
        status = grade.get("status")
        if status not in ("O", "X", "—", "미채점", "부분점수"):
            raise ValueError("채점 상태 오류")
        maximum = exam["points"][i]
        score = maximum if status == "O" else grade.get("score", 0) if status == "부분점수" else 0
        if not isinstance(score, (int, float)) or not 0 <= score <= maximum:
            raise ValueError(f"{i+1}번 부분점수 범위 오류")
        if status != "미채점":
            earned += score
            graded += 1
        key = f'{q["content"].get("unit", "미분류")} / Lv.{q["content"].get("level", "?")}'
        item = matrix.setdefault(key, {"earned": 0, "possible": 0, "pending": 0})
        if status == "미채점":
            item["pending"] += 1
        else:
            item["earned"] += score
            item["possible"] += maximum
    return {"exam": exam, "student": student, "grades": grades, "earned": earned,
            "possible": possible, "graded": graded, "matrix": matrix, "notes": notes or {},
            "expected_grade": expected_grade}


def report_html(report):
    student, exam = report["student"], report["exam"]
    school = "MIDDLE SCHOOL REPORT" if student.get("school") == "중등" else "HIGH SCHOOL REPORT"
    body = f'<div class="tag">{school}</div><h1>{text_html(student["name"])} · 학습 성적표</h1><p>{text_html(exam["title"])}</p>'
    label = "실제 점수" if report["graded"] == len(exam["questions"]) else "채점된 문항 점수 · 최종 점수 미산출"
    body += f'<p class="muted">{label}</p><div class="score">{report["earned"]} / {report["possible"]}</div>'
    if report.get("expected_grade"):
        body += '<p>교사 입력 예상등급: ' + text_html(report["expected_grade"]) + '</p>'
    body += '<h2>문항별 결과</h2><table><tr><th>번호</th><th>단원</th><th>난이도</th><th>결과</th></tr>'
    for i, (q, g) in enumerate(zip(exam["questions"], report["grades"]), 1):
        body += f'<tr><td>{i:02d}</td><td>{text_html(q["content"].get("unit", "미분류"))}</td><td>Lv.{q["content"].get("level", "?")}</td><td>{text_html(g["status"])}</td></tr>'
    body += '</table><h2>단원·난이도별 성취도</h2><table>'
    for key, item in report["matrix"].items():
        body += f'<tr><td>{text_html(key)}</td><td>{item["earned"]}/{item["possible"]}</td><td>미채점 {item["pending"]}</td></tr>'
    body += '</table><section class="report-page">'
    for key, title in (("weakness", "01 / 부족한 점과 보완 방법"), ("strength", "02 / 잘한 점과 강점"), ("summary", "TEACHER’S NOTE / 총평")):
        body += f'<h2>{title}</h2><p>{text_html(report["notes"].get(key, "작성 전"))}</p>'
    return page_html(body + '</section>', student["name"] + ' 성적표')
